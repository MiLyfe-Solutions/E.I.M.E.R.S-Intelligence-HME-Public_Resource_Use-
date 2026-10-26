from pathlib import Path
from openpyxl import Workbook,load_workbook
from fractions import Fraction
from datetime import datetime,timezone
import hashlib,json
D=Path.home()/'Desktop';p=D/'EIMERS_COLLATZ_HELD_COMPRESSION_V73.xlsx';rank=D/'EIMERS_COLLATZ_FRONTIER_PRIORITY_V70.xlsx';O=D/'EIMERS_COLLATZ_GATES6_7_V74.xlsx';J=O.with_suffix('.json')
if O.exists() or J.exists():raise SystemExit('HOLD_EXISTS')
w=load_workbook(p,read_only=True)
historical=[(int(r[0]),int(r[1]),int(r[2]),int(r[3]),str(r[4])) for r in list(w['02 RETAINED CERTS'].values)[1:]]
held=[(int(r[0]),int(r[1]),'V73_RETAINED') for r in list(w['03 COMPRESSED HELD'].values)[1:]]
ranks=list(load_workbook(rank,read_only=True)['01 PRIORITIES'].values)[1:]
assert len(historical)==30788 and len(held)==21200
def step(n):
 z=3*n+1;v=(z&-z).bit_length()-1;return z>>v,v
def cert(n,b,maxk=250):
 x=n;S=B=0
 for k in range(1,maxk+1):
  x,v=step(x);B=3*B+(1<<S);S+=v
  if S>=b:return None
  if x<n:return k,S,B
 return None
def vol(seq):return sum((Fraction(2)**(21-b) for n,b,*_ in seq),Fraction())
def canon(seq,M):
 spans=[]
 for n,b,*_ in seq:
  rev=int(format(n,'0%db'%b)[::-1],2);lo=rev<<(M-b);spans.append((lo,lo+(1<<(M-b))))
 spans.sort();out=[]
 for lo,hi in spans:
  assert not out or lo>=out[-1][1], 'OVERLAP'
  if out and lo==out[-1][1]:out[-1]=(out[-1][0],hi)
  else:out.append((lo,hi))
 return out
allnew=[];logs=[]
for deficit in (6,7):
 parents=[(int(r[0]),int(r[1])) for r in ranks if int(r[5])==deficit]
 assert len(parents)==(861 if deficit==6 else 699)
 existing={(n,b) for n,b,_ in held};assert set(parents).issubset(existing)
 active=parents[:];certs=[];pending=[];checks=0
 for depth in range(deficit+1):
  follow=[]
  for n,b in active:
   result=cert(n,b);checks+=1
   if result:certs.append((n,b,result[0],result[1],str(result[2])))
   elif depth<deficit:follow.extend(((n,b+1),(n+(1<<b),b+1)))
   else:pending.append((n,b,'NEXT_PRECISION_HOLD'))
  active=follow
 # MUST verify exact partition before advancing
 maxbits=max(b for n,b,*_ in certs+pending+parents)
 assert canon(certs+pending,maxbits)==canon(parents,maxbits)
 assert vol(certs)+vol(pending)==vol(parents)
 held=[r for r in held if (r[0],r[1]) not in set(parents)]+pending
 allnew.extend(certs)
 logs.append((deficit,len(parents),len(certs),len(pending),checks,str(vol(certs)),'PASS'))
 print('GATE',deficit,'PASS',len(certs),len(pending),flush=True)
m=max(b for n,b,*_ in historical+allnew+held)
original=[(int(r[0]),int(r[1])) for r in list(w['03 COMPRESSED HELD'].values)[1:]]
assert canon(allnew+held,m)==canon(original,m)
assert vol(historical)+vol(allnew)+vol(held)==59033
W=Workbook();s=W.active;s.title='01 GATES';s.append(['deficit','parents','new certificates','held children','checks','new equivalent 21bit','decision'])
for row in logs:s.append(row)
a=W.create_sheet('02 PRIOR CERTS');a.append(['r','bits','k','S','B'])
for row in historical:a.append(row)
b=W.create_sheet('03 NEW CERTS');b.append(['r','bits','k','S','B'])
for row in allnew:b.append(row)
c=W.create_sheet('04 HELD');c.append(['r','bits','reason'])
for row in held:c.append(row)
e=W.create_sheet('05 PROOF BOUNDARY')
for row in [('Source','v73 exact compressed frontier'),('Certification','Exact affine formula and S<b strict guard'),('All previous certificates','Retained unchanged'),('Full partition','Exact dyadic interval reconciliation'),('Global conjecture','HOLD_UNPROVEN')]:e.append(row)
W.save(O)
result={'utc':datetime.now(timezone.utc).isoformat(),'gates':logs,'retained_certificates':len(historical),'new_certificates':len(allnew),'remaining_cylinders':len(held),'certified_21bit':str(vol(historical)+vol(allnew)),'held_21bit':str(vol(held)),'source_sha256':hashlib.sha256(p.read_bytes()).hexdigest(),'sha256':hashlib.sha256(O.read_bytes()).hexdigest(),'global_theorem':'HOLD_UNPROVEN'}
J.write_text(json.dumps(result,indent=2)+'\n');print('V74_DONE',result,flush=True)
