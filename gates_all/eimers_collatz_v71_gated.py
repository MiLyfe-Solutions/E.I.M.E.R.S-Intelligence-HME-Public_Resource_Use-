from pathlib import Path
from openpyxl import Workbook,load_workbook
from fractions import Fraction
import hashlib,json,datetime
D=Path.home()/'Desktop';src=D/'EIMERS_COLLATZ_FRONTIER_PRIORITY_V70.xlsx';out=D/'EIMERS_COLLATZ_GATED_REFINEMENT_V71.xlsx';rec=out.with_suffix('.json')
if out.exists() or rec.exists():raise SystemExit('HOLD_OUTPUT_EXISTS')
raw=list(load_workbook(src,read_only=True)['01 PRIORITIES'].values)[1:]
assert len(raw)==13170
def step(n):
 z=3*n+1;v=(z&-z).bit_length()-1;return z>>v,v
def certificate(n,b,cap=250):
 x=n;S=B=0
 for k in range(1,cap+1):
  x,v=step(x);B=3*B+(1<<S);S+=v
  if S>=b:return None
  if x<n:return (k,S,B)
 return None
def first(n,cap=250):
 x=n;S=0
 for k in range(1,cap+1):
  x,v=step(x);S+=v
  if x<n:return k,S
 return None
def intervals(items,M):
 out=[]
 for n,b,*_ in items:
  reverse=int(format(n,'0%db'%b)[::-1],2);lo=reverse<<(M-b);out.append((lo,lo+(1<<(M-b))))
 return sorted(out)
def canon(items,M):
 merged=[]
 for lo,hi in intervals(items,M):
  assert not merged or lo>=merged[-1][1],'OVERLAP'
  if merged and lo==merged[-1][1]:merged[-1]=(merged[-1][0],hi)
  else:merged.append((lo,hi))
 return merged
def volume(items):return sum((Fraction(2)**(21-b) for n,b,*_ in items),Fraction())
allbase=[(int(r[0]),int(r[1])) for r in raw]
hold=[(int(r[0]),int(r[1]),'UNSELECTED') for r in raw if int(r[5]) not in (1,2,3)]
cert=[];log=[]
for deficit in (1,2,3):
 parents=[(int(r[0]),int(r[1])) for r in raw if int(r[5])==deficit]
 stage_cert=[];stage_hold=[];checks=0
 for n,b in parents:
  frontier=[(n,b)]
  for layer in range(deficit):
   nextfront=[]
   for r,bits in frontier:
    proven=certificate(r,bits)
    if proven:
     stage_cert.append((r,bits,*proven));checks+=1
    else:
     nextfront.extend(((r,bits+1),(r+(1<<bits),bits+1)));checks+=1
   frontier=nextfront
  for r,bits in frontier:
   proven=certificate(r,bits);checks+=1
   if proven:stage_cert.append((r,bits,*proven))
   else:stage_hold.append((r,bits,'NEEDS_FURTHER_BITS'))
 M=max([b for n,b,*_ in stage_cert+stage_hold+[(r,b) for r,b in parents]],default=21)
 assert canon(stage_cert+stage_hold,M)==canon(parents,M)
 assert volume(stage_cert)+volume(stage_hold)==volume(parents)
 cert.extend(stage_cert);hold.extend(stage_hold)
 log.append((deficit,len(parents),len(stage_cert),len(stage_hold),str(volume(stage_cert)),str(volume(stage_hold)),checks,'PASS'))
 print('GATE',deficit,'PASS','new_cert',len(stage_cert),'held',len(stage_hold),flush=True)
M=max(b for n,b,*_ in cert+hold)
assert canon(cert+hold,M)==canon(allbase,M)
assert volume(cert)+volume(hold)==Fraction(7205429,128)
W=Workbook();a=W.active;a.title='01 GATED PROOF'
a.append(['Additional bits group','Parents','New certificates','Held children','certified 21bit equivalent','held 21bit equivalent','checks','gate'])
for row in log:a.append(row)
b=W.create_sheet('02 NEW CERTIFICATES');b.append(['r','bits','k','S','B'])
for row in cert:b.append((row[0],row[1],row[2],row[3],str(row[4])))
c=W.create_sheet('03 NEXT HELD');c.append(['r','bits','reason'])
for row in hold:c.append(row)
e=W.create_sheet('04 GOVERNANCE')
for r in [('Baseline','v70 unresolved ranking, v69 immutable master'),('Steps','Gate 1-bit then 2-bit then 3-bit; every gate reconciled before next'),('Exact proof','S_k<b and A^k(r)<r'),('Coverage','Exact dyadic intervals, no full enumeration'),('Scope','Only v70 held classes; other v69 certificates untouched'),('Universal theorem','HOLD_UNPROVEN')]:e.append(r)
W.save(out)
j={'utc':datetime.datetime.now(datetime.timezone.utc).isoformat(),'gates':log,'new_certificates':len(cert),'new_equivalents':str(volume(cert)),'held_cylinders':len(hold),'held_equivalents':str(volume(hold)),'source_sha256':hashlib.sha256(src.read_bytes()).hexdigest(),'sha256':hashlib.sha256(out.read_bytes()).hexdigest(),'theorem':'HOLD_UNPROVEN'}
rec.write_text(json.dumps(j,indent=2)+'\n');print('V71_DONE',j,flush=True)
